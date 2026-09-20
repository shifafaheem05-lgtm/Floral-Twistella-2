#!/usr/bin/env python3
"""
post_to_facebook_instagram.py  -  Floral Twistella daily auto-poster
====================================================================
Publishes ONE post per run (photo or Reel) to a Facebook Page and an Instagram
Business account via the Meta Graph API.

* Content comes from content_calendar.xlsx (sheet "Content Calendar") - no dates, just Day 1, 2, 3...
* Progress is remembered in posted_state.json, so the calendar simply moves to the next day
  each time a post succeeds (and loops back to Day 1 after the last day, if enabled).
* Photos -> Facebook photo + Instagram image post.
  Reels  -> Facebook video + Instagram Reel.

Environment variables (GitHub: Settings -> Secrets and variables -> Actions):
    FB_PAGE_ID              Facebook Page ID                                  (secret)
    FB_PAGE_ACCESS_TOKEN    Long-lived PAGE access token                      (secret)
    IG_USER_ID              Instagram Business/Creator account ID             (secret)
  optional:
    MEDIA_BASE_URL          Public URL of the media/ folder. Default: this repo's raw.githubusercontent.com
                            URL (works for PUBLIC repos only). Instagram must be able to download the file.
    POST_TIMEZONE           Overrides the Settings sheet (default Asia/Karachi)
    GRAPH_API_VERSION       Default v26.0
    POST_DAY / DRY_RUN / FORCE_POST / POST_ONLY    Same as the --day / --dry-run / --force / --only flags (GitHub Actions)

Usage:
    python post_to_facebook_instagram.py --check          # validate calendar + media, post nothing
    python post_to_facebook_instagram.py --dry-run        # show what would be posted next
    python post_to_facebook_instagram.py                  # post the next day in the sequence
    python post_to_facebook_instagram.py --day 5 --dry-run
    python post_to_facebook_instagram.py --day 5 --only instagram   # manual re-post (state is NOT changed)
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
CALENDAR_FILE = ROOT / "content_calendar.xlsx"
STATE_FILE = ROOT / "posted_state.json"
SHEET_POSTS = "Content Calendar"
SHEET_SETTINGS = "Settings"

GRAPH_VERSION = os.getenv("GRAPH_API_VERSION", "v26.0")
GRAPH_URL = f"https://graph.facebook.com/{GRAPH_VERSION}"
REQUEST_TIMEOUT = 300
IG_CAPTION_LIMIT = 2200
IG_MAX_HASHTAGS = 30

# Column headers in the Excel sheet
C_DAY, C_TYPE, C_TITLE = "Day", "Post type", "Title"
C_MEDIA, C_IG, C_FB, C_ALT = "Media file", "Instagram caption", "Facebook caption", "Alt text / video description"


class PostError(RuntimeError):
    """A Graph API call or a precondition failed."""


# --------------------------------------------------------------------------- #
# Calendar + state
# --------------------------------------------------------------------------- #
def load_calendar(path: Path = CALENDAR_FILE) -> tuple[list[dict], dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_POSTS]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    missing = [h for h in (C_DAY, C_TYPE, C_MEDIA, C_IG, C_FB) if h not in headers]
    if missing:
        raise PostError(f"Calendar sheet is missing column(s): {', '.join(missing)}")

    posts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec.get(C_DAY) in (None, ""):
            continue
        rec[C_DAY] = int(rec[C_DAY])
        for key in (C_TYPE, C_TITLE, C_MEDIA, C_IG, C_FB, C_ALT):
            rec[key] = (str(rec.get(key) or "")).strip()
        posts.append(rec)
    posts.sort(key=lambda r: r[C_DAY])

    settings = {}
    if SHEET_SETTINGS in wb.sheetnames:
        for label, value, *_ in wb[SHEET_SETTINGS].iter_rows(min_row=2, values_only=True):
            if label:
                settings[str(label).strip()] = value
    return posts, settings


def load_state() -> dict:
    if STATE_FILE.exists():
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    return {"last_day": 0, "last_posted_date": None, "history": []}


def save_state(state: dict) -> None:
    state["history"] = state.get("history", [])[-90:]
    STATE_FILE.write_text(json.dumps(state, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def is_loop(settings: dict) -> bool:
    return str(settings.get("Loop after last day", "Yes")).strip().lower() in ("yes", "y", "true", "1")


def pick_next(posts: list[dict], state: dict, loop: bool) -> dict | None:
    nxt = int(state.get("last_day", 0)) + 1
    days = [p[C_DAY] for p in posts]
    if nxt > max(days):
        if not loop:
            print("Calendar finished and 'Loop after last day' is No - nothing to post.")
            return None
        nxt = min(days)
    return next(p for p in posts if p[C_DAY] == nxt)


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def validate(posts: list[dict]) -> list[str]:
    problems, seen = [], set()
    for p in posts:
        d, tag = p[C_DAY], f"Day {p[C_DAY]}"
        if d in seen:
            problems.append(f"{tag}: duplicate day number")
        seen.add(d)
        media = ROOT / p[C_MEDIA]
        if p[C_TYPE] not in ("Photo", "Reel"):
            problems.append(f"{tag}: Post type must be Photo or Reel (got '{p[C_TYPE]}')")
        if not media.exists():
            problems.append(f"{tag}: media file not found: {p[C_MEDIA]}")
        elif p[C_TYPE] == "Photo" and media.suffix.lower() not in (".jpg", ".jpeg"):
            problems.append(f"{tag}: Photo posts need a .jpg file")
        elif p[C_TYPE] == "Reel" and media.suffix.lower() != ".mp4":
            problems.append(f"{tag}: Reel posts need an .mp4 file")
        if not p[C_IG]:
            problems.append(f"{tag}: empty Instagram caption")
        if len(p[C_IG]) > IG_CAPTION_LIMIT:
            problems.append(f"{tag}: Instagram caption is {len(p[C_IG])} chars (limit {IG_CAPTION_LIMIT})")
        if p[C_IG].count("#") > IG_MAX_HASHTAGS:
            problems.append(f"{tag}: more than {IG_MAX_HASHTAGS} hashtags")
        if not p[C_FB]:
            problems.append(f"{tag}: empty Facebook caption")
    expected = list(range(1, len(posts) + 1))
    if sorted(seen) != expected:
        problems.append("Day numbers should run 1..N without gaps")
    return problems


# --------------------------------------------------------------------------- #
# Graph API helpers
# --------------------------------------------------------------------------- #
def graph_request(method: str, path: str, *, data=None, params=None, files=None, retries: int = 3) -> dict:
    url = f"{GRAPH_URL}/{path.lstrip('/')}"
    last_error = "unknown error"
    for attempt in range(1, retries + 1):
        try:
            resp = requests.request(method, url, data=data, params=params, files=files, timeout=REQUEST_TIMEOUT)
        except requests.RequestException as exc:
            last_error = f"network error: {exc}"
            if attempt < retries:
                time.sleep(4 * attempt)
                continue
            break
        try:
            payload = resp.json()
        except ValueError:
            payload = {"raw": resp.text[:500]}
        if resp.ok and "error" not in payload:
            return payload
        err = payload.get("error", payload)
        last_error = f"HTTP {resp.status_code}: {err}"
        transient = resp.status_code >= 500 or (isinstance(err, dict) and err.get("is_transient"))
        if transient and attempt < retries:
            time.sleep(5 * attempt)
            continue
        break
    raise PostError(f"{method} {path} failed -> {last_error}")


def media_url(media_file: str) -> str:
    """Public URL Instagram will download from."""
    name = media_file.replace("\\", "/").split("/")[-1]
    base = os.getenv("MEDIA_BASE_URL")
    if base:
        return f"{base.rstrip('/')}/{name}"
    repo = os.getenv("GITHUB_REPOSITORY")
    if repo:
        ref = os.getenv("GITHUB_REF_NAME") or "main"
        return f"https://raw.githubusercontent.com/{repo}/{ref}/media/{name}"
    raise PostError("Set MEDIA_BASE_URL (public URL of the media folder) when running outside GitHub Actions.")


def check_public_url(url: str) -> None:
    try:
        resp = requests.get(url, timeout=60, stream=True)
    except requests.RequestException as exc:
        raise PostError(f"Could not reach {url}: {exc}") from exc
    status = resp.status_code
    resp.close()
    if status != 200:
        raise PostError(
            f"Media URL is not publicly reachable (HTTP {status}): {url}\n"
            "Make sure media/ is pushed to a PUBLIC repo, or set MEDIA_BASE_URL to another public host."
        )


# --------------------------------------------------------------------------- #
# Facebook
# --------------------------------------------------------------------------- #
def post_facebook(post: dict, page_id: str, token: str) -> str:
    path = ROOT / post[C_MEDIA]
    if post[C_TYPE] == "Reel":
        files = {"source": (path.name, path.read_bytes(), "video/mp4")}
        data = {"description": post[C_FB], "title": post.get(C_TITLE, "")[:100], "published": "true",
                "access_token": token}
        res = graph_request("POST", f"{page_id}/videos", data=data, files=files)
        return str(res.get("id"))

    files = {"source": (path.name, path.read_bytes(), "image/jpeg")}
    data = {"caption": post[C_FB], "published": "true", "access_token": token}
    if post[C_ALT]:
        data["alt_text_custom"] = post[C_ALT]
    try:
        res = graph_request("POST", f"{page_id}/photos", data=data, files=files)
    except PostError as exc:
        if post[C_ALT] and "alt_text" in str(exc):  # alt text field rejected -> retry without it
            data.pop("alt_text_custom", None)
            res = graph_request("POST", f"{page_id}/photos", data=data, files=files)
        else:
            raise
    return str(res.get("post_id") or res.get("id"))


# --------------------------------------------------------------------------- #
# Instagram
# --------------------------------------------------------------------------- #
def _wait_for_container(container_id: str, token: str, tries: int, delay: int) -> None:
    for _ in range(tries):
        info = graph_request("GET", container_id, params={"fields": "status_code,status", "access_token": token})
        status = info.get("status_code")
        if status == "FINISHED":
            return
        if status in ("ERROR", "EXPIRED"):
            raise PostError(f"Instagram media container {status}: {info.get('status')}")
        time.sleep(delay)
    raise PostError("Timed out waiting for Instagram to process the media.")


def post_instagram(post: dict, ig_user_id: str, token: str) -> str:
    url = media_url(post[C_MEDIA])
    check_public_url(url)

    if post[C_TYPE] == "Reel":
        payload = {"media_type": "REELS", "video_url": url, "caption": post[C_IG],
                   "share_to_feed": "true", "access_token": token}
        wait = dict(tries=40, delay=6)
    else:
        payload = {"image_url": url, "caption": post[C_IG], "access_token": token}
        if post[C_ALT]:
            payload["alt_text"] = post[C_ALT]
        wait = dict(tries=12, delay=5)

    try:
        container = graph_request("POST", f"{ig_user_id}/media", data=payload)
    except PostError as exc:
        if "alt_text" in payload and "alt_text" in str(exc):
            payload.pop("alt_text")
            container = graph_request("POST", f"{ig_user_id}/media", data=payload)
        else:
            raise
    cid = container["id"]
    _wait_for_container(cid, token, **wait)
    res = graph_request("POST", f"{ig_user_id}/media_publish", data={"creation_id": cid, "access_token": token})
    return str(res.get("id"))


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def env_flag(name: str) -> bool:
    return os.getenv(name, "").strip().lower() in ("1", "true", "yes")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Floral Twistella daily Facebook + Instagram poster")
    ap.add_argument("--day", type=int, default=int(os.getenv("POST_DAY") or 0) or None,
                    help="Post a specific calendar day (manual run - progress is not changed)")
    ap.add_argument("--dry-run", action="store_true", default=env_flag("DRY_RUN"), help="Show the post, publish nothing")
    ap.add_argument("--only", choices=["both", "facebook", "instagram"], default=os.getenv("POST_ONLY") or "both")
    ap.add_argument("--force", action="store_true", default=env_flag("FORCE_POST"),
                    help="Post even if something was already posted today")
    ap.add_argument("--check", action="store_true", help="Validate calendar + media files and exit")
    args = ap.parse_args(argv)

    try:
        posts, settings = load_calendar()
    except (PostError, FileNotFoundError, KeyError) as exc:
        print(f"ERROR reading calendar: {exc}")
        return 1

    problems = validate(posts)
    if args.check:
        photos = sum(p[C_TYPE] == "Photo" for p in posts)
        print(f"{len(posts)} posts ({photos} photos, {len(posts) - photos} reels)")
        for prob in problems:
            print(f"  PROBLEM: {prob}")
        print("Calendar OK" if not problems else f"{len(problems)} problem(s) found")
        return 1 if problems else 0

    tz = ZoneInfo(os.getenv("POST_TIMEZONE") or str(settings.get("Timezone") or "Asia/Karachi"))
    today = datetime.now(tz).date()
    state = load_state()
    manual = bool(args.day) or args.only != "both"

    if args.day:
        post = next((p for p in posts if p[C_DAY] == args.day), None)
        if not post:
            print(f"ERROR: Day {args.day} is not in the calendar.")
            return 1
    else:
        if state.get("last_posted_date") == today.isoformat() and not args.force and not args.dry_run:
            print(f"Already posted today ({today}, Day {state['last_day']}). Use --force to post again.")
            return 0
        post = pick_next(posts, state, is_loop(settings))
        if post is None:
            return 0

    bad = [p for p in problems if p.startswith(f"Day {post[C_DAY]}:")]
    if bad:
        for prob in bad:
            print(f"ERROR: {prob}")
        return 1

    print(f"Date {today} | Day {post[C_DAY]}/{len(posts)} | {post[C_TYPE]} | {post[C_TITLE]}")
    print(f"Media: {post[C_MEDIA]}")

    if args.dry_run:
        print("\n----- INSTAGRAM CAPTION -----\n" + post[C_IG])
        print("\n----- FACEBOOK CAPTION -----\n" + post[C_FB])
        print("\n----- ALT TEXT -----\n" + post[C_ALT])
        print("\nDRY RUN: nothing was published and progress was not changed.")
        return 0

    results: dict[str, str] = {}
    failures = 0

    if args.only in ("both", "facebook"):
        try:
            post_id = post_facebook(post, os.environ["FB_PAGE_ID"], os.environ["FB_PAGE_ACCESS_TOKEN"])
            results["facebook"] = "ok"
            print(f"Facebook: published (id {post_id})")
        except KeyError as exc:
            results["facebook"] = f"missing env {exc}"
            print(f"Facebook: missing environment variable {exc}")
            failures += 1
        except PostError as exc:
            results["facebook"] = "failed"
            print(f"Facebook: FAILED -> {exc}")
            print(f"::error title=Facebook post failed::Day {post[C_DAY]}")
            failures += 1

    if args.only in ("both", "instagram"):
        try:
            post_id = post_instagram(post, os.environ["IG_USER_ID"], os.environ["FB_PAGE_ACCESS_TOKEN"])
            results["instagram"] = "ok"
            print(f"Instagram: published (id {post_id})")
        except KeyError as exc:
            results["instagram"] = f"missing env {exc}"
            print(f"Instagram: missing environment variable {exc}")
            failures += 1
        except PostError as exc:
            results["instagram"] = "failed"
            print(f"Instagram: FAILED -> {exc}")
            print(f"::error title=Instagram post failed::Day {post[C_DAY]}")
            failures += 1

    # Move the calendar forward only for scheduled runs where at least one platform succeeded
    if not manual and any(v == "ok" for v in results.values()):
        state["last_day"] = post[C_DAY]
        state["last_posted_date"] = today.isoformat()
        state.setdefault("history", []).append({"date": today.isoformat(), "day": post[C_DAY], **results})
        save_state(state)
        print(f"Progress saved: last_day = {post[C_DAY]}")

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
